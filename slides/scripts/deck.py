"""Helpers for building meeting slides from the deck template in files/ (keeps its theme and layouts).

Personal tooling for presentations, not part of the experiment code.
"""
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR

ROOT = Path(__file__).resolve().parents[2]  # slides/scripts/deck.py -> repo root
TEMPLATE = ROOT / "files/AI Alignemnt Davide all slides.pptx"

# positions taken from the template's own slides (EMU)
LEFT, WIDTH = 311700, 8520600
BODY_TOP = 1152475


def set_text(shape, text, level=0, new_paragraph=False):
    """Write text in the theme's black (dk1), like the template's own slides; the layout default is gray."""
    tf = shape.text_frame
    p = tf.add_paragraph() if new_paragraph else tf.paragraphs[0]
    p.text, p.level = text, level
    for run in p.runs:
        run.font.color.theme_color = MSO_THEME_COLOR.DARK_1


def new_deck(template=TEMPLATE):
    """The template with all its slides removed: same theme and layouts, empty deck."""
    prs = Presentation(template)
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
    return prs


def title_slide(prs, title, date):
    s = prs.slides.add_slide(prs.slide_layouts[0])  # TITLE
    set_text(s.placeholders[0], title)
    set_text(s.placeholders[1], date)
    return s


def slide(prs, title, bullets, figure=None, layout="below", body_height=1050000):
    """Title + bullets, optionally with a figure.

    bullets: list of str, or (level, str) for sub-bullets.
    figure: image path relative to the repo root.
    layout: "below" = text on top, figure under it (keep bullets to one line each);
            "right" = text in a left column, figure on the right.
    body_height: height of the text box for layout="below"; the figure fills the space under it.
    """
    s = prs.slides.add_slide(prs.slide_layouts[2])  # TITLE_AND_BODY
    set_text(s.placeholders[0], title)
    body = s.placeholders[1]
    for i, b in enumerate(bullets):
        level, text = b if isinstance(b, tuple) else (0, b)
        set_text(body, text, level, new_paragraph=i > 0)
    if figure is None:
        return s

    path = str(ROOT / figure)
    if layout == "right":
        body.left, body.top, body.width, body.height = 278775, BODY_TOP, 3985200, 3416400
        pic = s.shapes.add_picture(path, 4436675, 0, width=4395625)
        pic.top = max(773526, (prs.slide_height - pic.height) // 2)
    else:
        body.left, body.top, body.width, body.height = LEFT, BODY_TOP, WIDTH, body_height
        top = BODY_TOP + body_height + 50000
        pic = s.shapes.add_picture(path, LEFT, top)
        scale = min(WIDTH / pic.width, (prs.slide_height - top - 120000) / pic.height)
        pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
        pic.left = (prs.slide_width - pic.width) // 2
    return s
